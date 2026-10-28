"""
Собирает Word-файл со всеми событиями, решениями и последствиями игры «1 из 22 миллионов»
прямо из js/events.js и js/world.js — документ всегда совпадает с игрой.

Запуск (из папки 1in22m):  python tools/make_doc.py
Нужен python-docx:          pip install python-docx
"""
import json
import pathlib
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "1_из_22_миллионов_все_события.docx"

ACCENT = RGBColor(0xC2, 0x3B, 0x5E)
DIM = RGBColor(0x6B, 0x6B, 0x73)
GOOD = RGBColor(0x1E, 0x7D, 0x4A)
BAD = RGBColor(0xB4, 0x23, 0x18)
STAT = {"sec": "защищённость", "hap": "счастье", "kno": "знания", "rel": "отношения", "tru": "доверие к взрослым", "money": "₽", "exp": "цифровой след"}
FOCUS = {"study": "📚 Учёба", "friends": "🫶 Друзья", "games": "🎮 Игры", "create": "🎨 Блог и творчество", "sport": "⚽ Спорт", "work": "💼 Подработка"}
FLAGS = {
    "pwdWeak": "простой пароль", "pwdShared": "пароль знает кто-то ещё", "pwdSame": "один пароль везде", "twofa": "вход по коду (второй фактор)",
    "notesPass": "пароли в заметках", "sso": "вход через соцсеть", "overshare": "много личного в открытом доступе", "geo": "геометки",
    "profileClosed": "профиль только для друзей", "stranger": "общение с незнакомцем", "cheats": "взломанные программы и читы",
    "joinedMock": "участвовал(а) в насмешках", "rules": "договор с родителями", "bypass": "обходит родительский контроль",
    "cardSaved": "карта сохранена в сервисах", "cardLimit": "лимит и уведомления по карте", "cardShared": "карта у друга",
    "drop": "проводил(а) чужие переводы", "photoSent": "отправил(а) личное фото", "simGiven": "сим-карта на свой паспорт для другого",
    "smsGiven": "продиктовал(а) код", "selfban": "самозапрет на кредиты", "noUpdates": "отключены обновления", "codeword": "семейное кодовое слово",
    "voicePublic": "голос в открытом доступе", "gmWarned": "бабушка предупреждена", "gmSaved": "сбросил(а) звонок мошенников бабушке",
    "momWarned": "мама знает про SMS-развод", "club": "ходит в компьютерный клуб", "alone": "в интернете без взрослых",
    "leaked": "пароли утекли", "hasGm": "бабушка живёт с вами", "sibYoung": "младший брат", "gmLost": "бабушка потеряла деньги",
    "pyrFriends": "позвал(а) друзей в пирамиду", "track": "учёба / работа", "spec": "специальность",
}
SAMPLE = {"name": "Саша", "friend": "Женя", "city": "Тула", "mates": "одноклассники", "Mates": "Одноклассники",
          "to_study": "в школу", "in_study": "школы", "in_study_short": "школа", "klass_gen": "класса", "class_size": "26",
          "social_in": "во ВКонтакте / в соцсети", "msg_in": "в мессенджере", "msg": "мессенджер", "device": "телефон",
          "spec1": "Сварочное производство", "spec2": "Графический дизайн"}


def load(path, name):
    s = path.read_text(encoding="utf-8")
    start = s.index("window." + name + " = ") + len("window." + name + " = ")
    end = s.index("];", start) + 1
    return json.loads(s[start:end])


def render(text):
    def gender(m):
        pre, a, b = m.group(1), m.group(2), m.group(3)
        if a == "":
            return f"{pre}({b})"
        return f"{pre}{a}/{b}"
    t = re.sub(r"([А-Яа-яЁё]*)\{([^{}|]*)\|([^{}|]*)\}", gender, text)
    return re.sub(r"\{(\w+)\}", lambda m: SAMPLE.get(m.group(1), m.group(0)), t)


def fx_text(fx):
    out = []
    for k, v in (fx or {}).items():
        if k == "money":
            out.append("все накопления" if v <= -999999 else f"{'+' if v > 0 else '−'}{abs(v):,} ₽".replace(",", " "))
        else:
            out.append(f"{'+' if v > 0 else '−'}{abs(v)} {STAT[k]}")
    return ", ".join(out)


def set_text(o):
    parts = []
    for k, v in (o.get("set") or {}).items():
        name = FLAGS.get(k, k)
        if k in ("track", "spec"):
            parts.append(f"{name}: {render(str(v))}")
        else:
            parts.append(name if v else f"снимает «{name}»")
    for k, v in (o.get("add") or {}).items():
        parts.append({"exp": "цифровой след", "click": "рискованный клик"}.get(k, k) + f" +{v}")
    return "; ".join(parts)


def para(doc, text="", bold=False, italic=False, color=None, size=None, style=None, after=4):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p


def labeled(doc, label, text, color=None, indent=0.0, bold_text=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(label + " ")
    r.bold = True
    if color is not None:
        r.font.color.rgb = color
    r2 = p.add_run(text)
    r2.bold = bold_text
    return p


def when(e):
    y = e["year"]
    years = "любые годы" if y[0] <= 1999 and y[1] >= 2100 else (f"с {y[0]} года" if y[1] >= 2100 else f"{y[0]}–{y[1]}")
    return f"Возраст {e['age'][0]}–{e['age'][1]} · {years}"


def main():
    events = load(ROOT / "js" / "events.js", "EVENTS")
    risks = load(ROOT / "js" / "events.js", "RISKS")
    by_id = {e["id"]: e for e in events}
    facts = dict(re.findall(r'(\w+):\s+"(.*?)",?\n', (ROOT / "js" / "world.js").read_text(encoding="utf-8").split("FACTS: {")[1]))

    # кто вызывает какое последствие
    callers = {}
    for e in events:
        for o in e.get("opts", []) + [e]:
            for l in o.get("later", []) or []:
                callers.setdefault(l["ev"], []).append(f"«{render(e.get('title', e['id']))}» → «{render(o.get('t', ''))}»" if o is not e else f"событие {e['id']}")
    for r in risks:
        callers.setdefault(r["ev"], []).append("ежегодный риск")
        if r.get("safe"):
            callers.setdefault(r["safe"], []).append("ежегодный риск, если есть защита")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)

    doc.add_heading("1 из 22 миллионов", 0)
    para(doc, "Симуляция взросления о цифровой безопасности — разминка перед рабочими листами «Мастерской безопасных решений». "
              "Все решения, варианты, последствия и правила.", color=DIM, size=12)
    para(doc, "Игра: https://irinilou.github.io/1in22m.life/", bold=True)

    doc.add_heading("Как устроена игра", 1)
    for line in [
        "По образцу «1 из 8 миллиардов». Участник выбирает мир (прошлое — рождение в 1990–1994, сегодня — 2008–2012, будущее — 2018–2022), "
        "способ выбора региона, длительность (≈5, ≈8 или ≈12 минут) и пол. Случай выбирает один из 89 регионов (с показом вероятности), "
        "населённый пункт, семью, достаток, родителей, бабушку, четырёх ровесников.",
        "Жизнь идёт с 7 до 23 лет сама: один год длится 6, 10 или 16 секунд. Есть пауза, скорость 1×–8× и кнопка «До следующего решения».",
        "Фокус года: учёба, друзья, игры, блог и творчество, спорт, подработка (с 14 лет). Фокус меняет шкалы каждый год и делает "
        "связанные с ним события вероятнее. Его можно сменить в любой момент.",
        "Решения всплывают окном: заголовок, ситуация, иногда блок «Данные» с фактом и источником, 3–4 пронумерованных варианта. "
        "Варианты написаны нейтрально, без подсказок. У решений нет цвета в летописи — красным и зелёным подсвечиваются только последствия.",
        "Последствия часто приходят не сразу: через минуту игрового года или через 1–2 года. В летописи у них подпись "
        "«↳ последствие: …», а в финале — список «Последствия, которые пришли позже».",
        "Риски. Раз в год игра проверяет, что накопилось: простой или общий пароль, читы, геометки, личное в открытом доступе, утечки, "
        "рискованные клики. Отсюда — взломы, заражения, травля, преследование. Защита (вход по коду, предупреждённая бабушка, "
        "кодовое слово) превращает атаку в «атаку отбили».",
        "Доверие к взрослым — главная шкала для педагогов. У части событий два текста: если доверие 55 и выше, герой рассказывает "
        "взрослым и всё заканчивается мягче; если ниже — молчит, и потери больше.",
        "Травля не стоит отдельной темой: она появляется как результат цифровых решений — слитые скриншоты, взломанная страница, "
        "старые посты, фото, отправленное «только для меня».",
        "Развилки пути: после 9 класса (10 класс или колледж с одной из двух специальностей), после 11 (вуз, колледж, работа), "
        "диплом колледжа (работа или учёба дальше). В 16 лет — «Пауза»: кому вы напишете, если в сети что-то пойдёт не так.",
        "Финал «Запись о жизни»: сколько решений, взломов, случаев травли, сколько денег потеряли вы и семья, сколько атак отбито, "
        "карта жизни, «Взгляд педагога» с выводами именно из этой жизни, все решения и вся летопись.",
        "Примерно 20 решений за жизнь в режиме ≈8 минут (из 56 возможных, плюс развилки и последствия). Набор зависит от года, "
        "возраста, фокуса и прошлых решений, поэтому две одинаковые жизни в зале практически невозможны.",
    ]:
        para(doc, line, style="List Bullet", after=2)

    doc.add_heading("Как читать этот файл", 1)
    para(doc, "Род дан через скобки и косую черту: «дал(а)», «доволен/довольна». Имя друга в примере — Женя, город — Тула. "
              "«Скрыто» — что вариант незаметно меняет в жизни героя; это влияет на будущие риски. «Позже» — отложенное последствие "
              "с вероятностью.")

    def write_decision(e, n=None):
        title = f"{n}. {render(e['title'])}" if n else render(e["title"])
        doc.add_heading(title, 3)
        meta = when(e)
        if e.get("focus"):
            meta += " · чаще при фокусе: " + ", ".join(FOCUS[f] for f in e["focus"])
        if e.get("need"):
            meta += " · только если: " + ", ".join(FLAGS.get(k, k) for k in e["need"])
        para(doc, meta, color=DIM, size=9.5, after=2)
        para(doc, render(e["text"]))
        if e.get("fact"):
            labeled(doc, "Данные:", facts.get(e["fact"], ""), color=DIM)
        for i, o in enumerate(e["opts"], 1):
            labeled(doc, f"{i}.", render(o["t"]), bold_text=True)
            labeled(doc, "Что происходит:", render(o["r"]), indent=0.6)
            ft = fx_text(o.get("fx"))
            if ft:
                labeled(doc, "Шкалы:", ft, color=DIM, indent=0.6)
            stx = set_text(o)
            if stx:
                labeled(doc, "Скрыто:", stx, color=DIM, indent=0.6)
            for l in o.get("later", []) or []:
                tgt = by_id.get(l["ev"], {})
                d = "в том же году" if l["in"][1] == 0 else f"через {l['in'][0]}–{l['in'][1]} г."
                labeled(doc, "Позже:", f"с вероятностью {int(l['p'] * 100)} % {d} — «{render(tgt.get('title') or tgt.get('text', l['ev'])[:70])}»",
                        color=ACCENT, indent=0.6)

    decisions = [e for e in events if e["kind"] == "decision" and not e.get("trigger") and not e.get("path")]
    order = sorted(decisions, key=lambda e: (e["year"][0] > 2015, e["year"][0], e["age"][0]))
    doc.add_page_break()
    doc.add_heading(f"Решения ({len(decisions)})", 1)
    para(doc, "Выпадают случайно с учётом возраста, года и фокуса. Сначала — решения, возможные в 2000-е, потом — более поздние.", color=DIM)
    for n, e in enumerate(order, 1):
        write_decision(e, n)

    doc.add_page_break()
    doc.add_heading("Развилки пути", 1)
    for e in [x for x in events if x.get("path")]:
        write_decision(e)

    doc.add_page_break()
    doc.add_heading("Последствия-решения", 1)
    para(doc, "Не выпадают сами — только как последствие прошлых решений или ежегодного риска.", color=DIM)
    for e in [x for x in events if x["kind"] == "decision" and x.get("trigger")]:
        write_decision(e)
        if callers.get(e["id"]):
            labeled(doc, "Откуда берётся:", "; ".join(sorted(set(callers[e["id"]]))), color=DIM)

    doc.add_heading("Последствия-события", 1)
    for e in [x for x in events if x["kind"] == "event" and x.get("trigger")]:
        color = BAD if e.get("tone") == "bad" else GOOD if e.get("tone") == "good" else None
        p = labeled(doc, e["id"] + ":", render(e["text"]), color=color)
        if e.get("hi"):
            labeled(doc, "Если доверие к взрослым ≥ 55:", render(e["hi"]), color=GOOD, indent=0.6)
        for al in e.get("alt", []) or []:
            labeled(doc, "Если " + ", ".join(FLAGS.get(k, k) for k in al["if"]) + ":", render(al["text"]), color=GOOD, indent=0.6)
        ft = fx_text(e.get("fx"))
        if ft:
            labeled(doc, "Шкалы:", ft, color=DIM, indent=0.6)
        if callers.get(e["id"]):
            labeled(doc, "Откуда берётся:", "; ".join(sorted(set(callers[e["id"]]))), color=DIM, indent=0.6)

    doc.add_heading("Ежегодные риски", 1)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Событие", "Когда", "Что повышает шанс", "Что защищает"]):
        t.rows[0].cells[i].text = h
    for r in risks:
        row = t.add_row().cells
        tgt = by_id[r["ev"]]
        row[0].text = r["ev"] + " — " + render(tgt.get("title") or tgt["text"][:60]) + ("…" if not tgt.get("title") else "")
        row[1].text = f"{r['age'][0]}–{r['age'][1]} лет, с {r['year'][0]}" + (" · один раз" if r.get("once") else "")
        ups = [f"{FLAGS.get(k, k)} +{int(v * 100)} %" for k, v in (r.get("flags") or {}).items()]
        ups += [f"{ {'exp': 'цифровой след', 'click': 'рискованный клик'}[k]} +{int(v * 100)} % за каждый" for k, v in (r.get("nums") or {}).items()]
        row[2].text = (f"основа {int(r['base'] * 100)} %; " if r.get("base") else "") + "; ".join(ups)
        row[3].text = ", ".join(FLAGS.get(k, k) for k in r.get("protect", [])) or "доверие к взрослым смягчает исход"
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    doc.add_heading("Фоновые события жизни", 1)
    for e in [x for x in events if x["kind"] == "event" and not x.get("trigger")]:
        labeled(doc, when(e) + ":", render(e["text"]), color=DIM)

    doc.add_heading("Данные (факты с источниками)", 1)
    for k, v in facts.items():
        para(doc, "• " + v, after=2)

    doc.save(OUT)
    print("Готово:", OUT)


if __name__ == "__main__":
    main()
